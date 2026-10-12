"""校验 Excel 底稿：公式齐不齐、重算值跟引擎对不对得上。

配合 check_excel.ps1 使用（后者负责用 Excel 打开、重算、存一份带缓存值的副本）：

    py valuation/tools/smoke_engine.py                      # 跑引擎，出底稿
    powershell -File valuation/tools/check_excel.ps1 -src valuation/demo/valuation_model.xlsx
    py valuation/tools/verify_workbook.py                   # 比对

按 B 列标签定位，不依赖固定行号——底稿加一行注释就会让硬编码行号失效。
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent

# (工作表, B 列标签, 结果 JSON 的取值路径)
CHECKS = [
    ("参数假设", "WACC", ("dcf", "wacc")),
    ("DCF测算", "每股价值", ("dcf", "per_share")),
    ("DCF测算", "企业价值 EV", ("dcf", "enterprise_value")),
    ("相对估值", "隐含每股价值", ("relative", "implied_per_share")),
    ("三法融合", "融合每股价值", ("blended", "blended_per_share")),
]

TOL = 1e-6


def dig(obj, path):
    for key in path:
        if not isinstance(obj, dict):
            return None
        obj = obj.get(key)
    return obj


def find_row(ws, label: str) -> int | None:
    for row in range(1, ws.max_row + 1):
        if ws.cell(row=row, column=2).value == label:
            return row
    return None


def main() -> int:
    recalc = ROOT / "demo" / "excel_check_recalc.xlsx"
    result_path = ROOT / "demo" / "output_sample.json"
    src = ROOT / "demo" / "valuation_model.xlsx"

    for path in (recalc, result_path, src):
        if not path.exists():
            print(f"✘ 找不到 {path}。先跑 smoke_engine.py 和 check_excel.ps1。")
            return 2

    # 公式数量：底稿必须真的是公式联动的，不是静态报告
    live = openpyxl.load_workbook(src, data_only=False)
    formula_count = sum(
        1 for ws in live.worksheets for row in ws.iter_rows() for c in row
        if isinstance(c.value, str) and c.value.startswith("=")
    )

    wb = openpyxl.load_workbook(recalc, data_only=True)
    result = json.loads(result_path.read_text(encoding="utf-8"))

    print(f"工作表 {len(wb.sheetnames)} 张：{'、'.join(wb.sheetnames)}")
    print(f"公式总数 {formula_count}")
    print(f"必需表：", end="")
    required = ["说明", "参数假设", "DCF测算", "相对估值", "分部估值",
                "三法融合", "敏感性分析", "数据来源"]
    missing = [n for n in required if n not in wb.sheetnames]
    print("齐全" if not missing else f"✘ 缺 {missing}")

    bad = 0
    print(f"\n{'工作表':<8}{'科目':<16}{'底稿':>18}{'引擎':>18}  判定")
    for sheet, label, path in CHECKS:
        if sheet not in wb.sheetnames:
            print(f"{sheet:<8}{label:<16}{'—':>18}{'—':>18}  ✘ 表不存在")
            bad += 1
            continue
        ws = wb[sheet]
        row = find_row(ws, label)
        if row is None:
            print(f"{sheet:<8}{label:<16}{'—':>18}{'—':>18}  ✘ 找不到该科目")
            bad += 1
            continue
        excel_v = ws.cell(row=row, column=3).value
        engine_v = dig(result, path)
        ok = (isinstance(excel_v, (int, float)) and isinstance(engine_v, (int, float))
              and abs(excel_v - engine_v) <= TOL * max(1.0, abs(engine_v)))
        if not ok:
            bad += 1
        print(f"{sheet:<8}{label:<16}{excel_v!s:>18}{engine_v!s:>18}  {'✓' if ok else '✘ 不一致'}")

    print()
    if missing:
        print("✘ 底稿缺表——「三种方法」必须都在底稿里")
        return 1
    if bad:
        print(f"✘ {bad} 项对不上：底稿重算值与引擎结论不一致")
        return 1
    if formula_count < 100:
        print(f"✘ 公式只有 {formula_count} 个，底稿多半退化成静态报告了")
        return 1
    print("✓ 底稿公式联动、八表齐全、关键值与引擎一致")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
